"""Top N 岗位 JD 详情补采 + 二次精排（M1 补充 + M6.1 增强）。

流程：
1. select_top() 从已精排岗位中按 fit_score 选 Top N（排除同公司重复）
2. fetch_details() 用 Playwright 逐条打开详情页，提取 JD 全文
3. rescore() 对拿到 JD 的岗位用 LLM 做第二轮评估（JD 全文 vs 简历）
4. 输出 APPLY_NOW / REVIEW / SKIP 投递候选清单

原则：
- 单条失败不中断整批
- 连续 3 次风控错误停止自动补采
- 每条最多尝试 2 次
- 不用自动化绕过验证码
"""

from __future__ import annotations

import csv
import json
import re
import time
import random
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from .config import DATA_DIR, REPORTS_DIR, RESUME_DIR, load_config
from . import db, llm, score as score_mod

ENRICH_DIR = DATA_DIR / "enriched"

# JD 详情页 DOM 选择器（Boss直聘 Web 版，按优先级尝试）
_JD_SELECTORS = [
    ".job-sec-text",
    ".job-detail-section",
    ".detail-section",
    ".job-detail",
    "div[class*='job-detail'] div[class*='text']",
    ".detail-content",
]

# 详情页加载等待上限
_PAGE_LOAD_TIMEOUT_MS = 15000
_JD_WAIT_TIMEOUT_MS = 8000

# 连续风控错误上限
_MAX_CONSECUTIVE_SECURITY = 3


# ============================================================
# Top N 选择
# ============================================================

def select_top(n: int = 30) -> list[dict]:
    """从已精排岗位中选 Top N，确保多样性（同公司不超过 2 条）。"""
    conn = db.connect()
    rows = conn.execute(
        """SELECT s.fit_score, s.verdict, s.archetype, s.seniority_ok, s.authenticity,
                  s.reasons, s.highlights, s.gaps,
                  j.id, j.job_id, j.url, j.title, j.company, j.company_size, j.industry,
                  j.salary_text, j.salary_min, j.salary_max, j.salary_months,
                  j.city, j.district, j.experience, j.degree, j.tags, j.jd_text
           FROM job_scores s JOIN jobs j ON j.id = s.job_pk
           WHERE s.seniority_ok = 1 AND s.authenticity != '低'
             AND j.url IS NOT NULL AND j.url != ''
           ORDER BY s.fit_score DESC, j.salary_max DESC"""
    ).fetchall()

    picked = []
    company_count: Counter = Counter()
    seen_ids = set()

    for r in rows:
        r = dict(r)
        jid = r["id"]
        if jid in seen_ids:
            continue
        comp = (r.get("company") or "").strip()
        if company_count.get(comp, 0) >= 2:
            continue  # 同公司最多 2 条
        seen_ids.add(jid)
        company_count[comp] += 1
        picked.append(r)
        if len(picked) >= n:
            break

    return picked


# ============================================================
# JD 详情补采
# ============================================================

@dataclass
class FetchResult:
    job_pk: int
    job_id: str
    url: str
    success: bool
    jd_text: str = ""
    failure_reason: str = ""


def _extract_jd_from_page(page) -> str:
    """从 Boss直聘 详情页 DOM 提取 JD 文本。"""
    for sel in _JD_SELECTORS:
        try:
            el = page.query_selector(sel)
            if el:
                text = el.inner_text()
                if text and len(text.strip()) > 30:
                    return text.strip()
        except Exception:
            continue

    # 兜底：尝试从整个页面提取可能包含 JD 的文本块
    try:
        body = page.inner_text("body")
        # 找常见 JD 段落标记
        patterns = [
            r"岗位职责[\s\S]{50,}?(?=任职要求|岗位要求|职位要求|$)",
            r"职位描述[\s\S]{50,}?(?=任职要求|岗位要求|职位要求|$)",
            r"工作内容[\s\S]{50,}?(?=任职要求|岗位要求|职位要求|$)",
        ]
        for pat in patterns:
            m = re.search(pat, body)
            if m:
                return m.group(0).strip()
    except Exception:
        pass

    return ""


def fetch_details(candidates: list[dict]) -> list[FetchResult]:
    """用 Playwright 逐条获取 Top N 岗位的详情页 JD（headless，复用登录 cookie）。"""
    from playwright.sync_api import sync_playwright
    from .collect_boss import (
        _session_path, _is_logged_in, _launch,
        _hit_security_check, _wait_if_security_check,
    )

    state = _session_path()
    if not state.exists():
        print("✗ 未找到登录态。请先执行: python -m job_copilot collect login")
        return [_fail(r, "无登录态") for r in candidates]

    results = []
    consecutive_security = 0
    cfg = load_config()
    lo = float(cfg.collect.get("min_interval_sec", 3))
    hi = float(cfg.collect.get("max_interval_sec", 8))

    with sync_playwright() as p:
        browser, context = _launch(p, state, headless=True)
        page = context.new_page()

        for i, r in enumerate(candidates):
            if consecutive_security >= _MAX_CONSECUTIVE_SECURITY:
                print(f"⚠ 连续 {consecutive_security} 次风控错误，停止自动补采。"
                      f"剩余 {len(candidates) - i} 条将标记为 pending。")
                for remaining in candidates[i:]:
                    results.append(_fail(remaining, "风控停止（批量）"))
                break

            url = (r.get("url") or "").strip()
            if not url:
                results.append(_fail(r, "无详情页 URL"))
                continue

            jd_text = ""
            failure = ""
            for attempt in range(2):
                try:
                    page.goto(url, wait_until="domcontentloaded",
                              timeout=_PAGE_LOAD_TIMEOUT_MS)
                    page.wait_for_timeout(2000)

                    if _hit_security_check(page):
                        consecutive_security += 1
                        failure = f"风控验证页（第{attempt + 1}次）"
                        _wait_if_security_check(page, timeout_sec=60)
                        continue

                    # 尝试等待 JD 内容出现
                    try:
                        page.wait_for_selector(
                            ",".join(_JD_SELECTORS),
                            timeout=_JD_WAIT_TIMEOUT_MS,
                        )
                    except Exception:
                        pass  # 选择器可能不匹配，继续用兜底提取

                    page.wait_for_timeout(1000)
                    jd_text = _extract_jd_from_page(page)
                    if jd_text:
                        consecutive_security = 0
                        failure = ""
                        break
                    else:
                        failure = f"未提取到 JD 文本（第{attempt + 1}次）"
                except Exception as e:
                    failure = f"页面加载失败: {type(e).__name__}"

            if jd_text:
                results.append(FetchResult(
                    job_pk=r["id"], job_id=r.get("job_id", ""),
                    url=url, success=True, jd_text=jd_text,
                ))
                print(f"  [{i+1}/{len(candidates)}] ✓ {r.get('title','')[:24]} "
                      f"({len(jd_text)}字)")
            else:
                results.append(FetchResult(
                    job_pk=r["id"], job_id=r.get("job_id", ""),
                    url=url, success=False, failure_reason=failure,
                ))
                print(f"  [{i+1}/{len(candidates)}] ✗ {r.get('title','')[:24]} "
                      f"— {failure}")

            # 保存已获取的 JD 到数据库
            if jd_text:
                conn = db.connect()
                conn.execute(
                    "UPDATE jobs SET jd_text = ? WHERE id = ?",
                    (jd_text, r["id"]),
                )
                conn.commit()

            # 间隔延迟
            if i < len(candidates) - 1:
                time.sleep(random.uniform(lo, hi))

        try:
            browser.close()
        except Exception:
            pass

    return results


def _fail(r: dict, reason: str) -> FetchResult:
    return FetchResult(
        job_pk=r["id"], job_id=r.get("job_id", ""),
        url=r.get("url", ""), success=False, failure_reason=reason,
    )


# ============================================================
# 二次精排（含 JD 全文）
# ============================================================

ENRICHED_SYSTEM = """你是资深技术招聘顾问，帮一位【AI应用/Agent工程】转行候选人评估岗位匹配度。

这次你拥有**完整 JD 全文**，请基于 JD 原文与候选人简历做精确对比。
严格基于简历事实，不得臆造经历。
只输出 JSON，字段：
- fit_score: 1-5 匹配度（基于 JD 全文与简历的综合判断）
  5=高度契合优先投；4=契合可投；3=模棱两可/利弊相当；2=偏不合适；1=基本不匹配。
- verdict: "投" | "慎投" | "不投"（须与 fit_score 一致：≥4→投, =3→慎投, ≤2→不投）
- archetype: 该岗真实画像（AI应用工程/Agent工程/大模型算法研究/后端工程/前端全栈/数据工程/其它）
- seniority_ok: true/false
- authenticity: "高"|"中"|"低"（岗位真实性）
- matched_evidence: 2-4条，JD中明确要求且简历能证明的具体点（每条引用JD原文片段+简历对应证据）
- missing_requirements: 2-4条，JD硬性要求但简历不满足的点
- hard_blockers: 硬性阻塞项（如"要求5年Java经验但候选人无Java经历"），无则为空数组
- acceptable_gaps: 可接受缺口（要求但可通过短期学习/项目补齐），无则为空数组
- reasons: 3条以内简短理由
- recommendation: "APPLY_NOW" | "REVIEW" | "SKIP"
  APPLY_NOW=高度契合可立即投递
  REVIEW=有价值但需人工确认某些点
  SKIP=不建议投递
- confidence: "high"|"medium"|"low"（本次判断的置信度，基于JD完整度+信息对称度）
- highlights: 简历应突出的2-3个真实点
- score_change_note: 一句话说明与只看标签时的评分差异（如"JD显示需要5年后端经验，之前仅看标签高估"）

注意：不要因为缺一项非硬性技能就判定不匹配；区分"硬性门槛"和"加分项"。"""

ENRICHED_USER_TMPL = """【候选人简历（唯一事实依据）】
{resume}

【岗位完整信息】
职位：{title}
公司：{company}（{industry}，{size}）
薪资：{salary}　经验要求：{exp}　学历：{degree}　城市：{city}
技能标签：{tags}

【JD 全文】
{jd_text}

请基于以上完整 JD 输出该岗的 JSON 评估。特别注意：
- 引用 JD 原文具体描述与简历对应证据做匹配判断
- 区分"硬性门槛"（如学历/经验年限/必须会某语言）和"加分项"（如"了解XX更佳"）
- 对转行候选人，关注实际能力而非毕业年限"""


def _parse_tags(r: dict) -> list:
    try:
        return json.loads(r.get("tags") or "[]")
    except Exception:
        return []


def rescore_one(conn, r: dict, resume: str, model: str) -> dict:
    """用完整 JD 做二次精排。"""
    jd_text = (r.get("jd_text") or "").strip()
    if not jd_text:
        # 无 JD 全文 → PENDING_JD（等待补采），保留原始分供排序参考
        existing = db.load_scores(conn).get(r["id"], {})
        if existing:
            existing["recommendation"] = "PENDING_JD"
            existing["confidence"] = "low"
            existing["score_change_note"] = "等待JD全文验证——当前评分仅基于列表页标签"
            existing["matched_evidence"] = existing.get("highlights", [])
            existing["missing_requirements"] = existing.get("gaps", [])
            existing["hard_blockers"] = []
            existing["acceptable_gaps"] = []
            return existing
        return {"fit_score": 0, "verdict": "慎投", "recommendation": "PENDING_JD",
                "confidence": "low", "score_change_note": "无JD全文且无历史精排"}

    # 截断过长 JD（保留职责和要求部分）
    if len(jd_text) > 2500:
        jd_text = jd_text[:2500] + "…"

    user = ENRICHED_USER_TMPL.format(
        resume=resume[:3500],
        title=r.get("title") or "", company=r.get("company") or "",
        industry=r.get("industry") or "", size=r.get("company_size") or "",
        salary=r.get("salary_text") or "", exp=r.get("experience") or "",
        degree=r.get("degree") or "", city=r.get("city") or "",
        tags="、".join(_parse_tags(r)),
        jd_text=jd_text,
    )
    try:
        s = llm.chat_json(ENRICHED_SYSTEM, user, model=model, max_tokens=3000)
    except Exception as e:
        return {"fit_score": 0, "verdict": "慎投", "recommendation": "REVIEW",
                "confidence": "low", "score_change_note": f"LLM调用失败: {e}"}

    s["verdict"] = score_mod._verdict_from_fit(s.get("fit_score"))
    s.setdefault("recommendation", _rec_from_fit(s.get("fit_score")))
    s.setdefault("confidence", "medium")
    s.setdefault("score_change_note", "")
    s.setdefault("matched_evidence", [])
    s.setdefault("missing_requirements", [])
    s.setdefault("hard_blockers", [])
    s.setdefault("acceptable_gaps", [])

    # 重新保存精排结果（覆盖第一轮）
    db.save_score(conn, r["id"], s, model)
    return s


def _rec_from_fit(fit) -> str:
    try:
        f = float(fit)
    except (TypeError, ValueError):
        return "REVIEW"
    if f >= 4:
        return "APPLY_NOW"
    if f <= 2:
        return "SKIP"
    return "REVIEW"


def rescore_all(candidates: list[dict], fetch_results: list[FetchResult]) -> list[dict]:
    """对所有候选做二次精排（有 JD 的用全文，无 JD 的保留原分）。"""
    conn = db.connect()
    resume = score_mod._resume_text()
    if not resume:
        print("✗ 未找到母版简历，无法精排。")
        return []

    cfg = load_config()
    model = cfg.llm["model_analysis"]

    # fetch result lookup
    fr_map = {fr.job_pk: fr for fr in fetch_results}

    enriched = []
    print(f"→ 二次精排 {len(candidates)} 个候选（含JD全文比对）…")
    for i, r in enumerate(candidates):
        # 用 DB 中最新 jd_text（fetch_details 已写入）
        row = dict(conn.execute("SELECT * FROM jobs WHERE id = ?", (r["id"],)).fetchone())
        try:
            s = rescore_one(conn, row, resume, model)
            fr = fr_map.get(r["id"])
            s["_job"] = row
            s["_fetch_success"] = fr.success if fr else False
            s["_jd_length"] = len(row.get("jd_text") or "")
            enriched.append(s)
            print(f"  [{i+1}/{len(candidates)}] {row.get('title','')[:24]} → "
                  f"fit={s.get('fit_score')} {s.get('recommendation','?')} "
                  f"(JD:{s['_jd_length']}字, conf={s.get('confidence','?')})")
        except Exception as e:
            print(f"  [{i+1}/{len(candidates)}] {row.get('title','')[:24]} → "
                  f"评分失败: {type(e).__name__}")

    # 按 recommendation 分组排序：APPLY_NOW → REVIEW → SKIP
    # 同组内按 fit_score 降序
    rec_order = {"APPLY_NOW": 0, "REVIEW": 1, "SKIP": 2}
    enriched.sort(key=lambda x: (
        rec_order.get(x.get("recommendation", "REVIEW"), 9),
        -(x.get("fit_score") or 0),
    ))
    return enriched


# ============================================================
# 报告生成
# ============================================================

def _write_reports(enriched: list[dict], candidates: list[dict],
                   fetch_results: list[FetchResult]) -> tuple[Path, Path, Path]:
    """生成 Markdown + CSV + Excel 投递候选清单。"""
    ENRICH_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    # 统计
    fr_success = sum(1 for fr in fetch_results if fr.success)
    rec_counts = Counter(s.get("recommendation") for s in enriched)
    for k in ["APPLY_NOW", "REVIEW", "SKIP", "PENDING_JD"]:
        rec_counts.setdefault(k, 0)

    # --- Markdown ---
    today = date.today().isoformat()
    # 标题区分：有 JD 精排 vs 待补采
    has_real_jd = rec_counts['APPLY_NOW'] + rec_counts['REVIEW'] + rec_counts['SKIP'] > 0
    title = "投递候选清单（JD二次精排）" if has_real_jd else "投递候选清单（待补充JD）"
    L = [
        f"# {title}",
        "",
        f"> 生成：{today} ｜ 候选：{len(candidates)} 个 ｜ JD 补采成功：{fr_success} 个",
        f"> APPLY_NOW {rec_counts['APPLY_NOW']} · REVIEW {rec_counts['REVIEW']} · SKIP {rec_counts['SKIP']} · PENDING_JD {rec_counts['PENDING_JD']}",
        "",
        "---",
        "",
    ]

    current_group = None
    for i, s in enumerate(enriched):
        rec = s.get("recommendation", "PENDING_JD")
        if rec != current_group:
            current_group = rec
            emoji = {"APPLY_NOW": "🟢", "REVIEW": "🟡", "SKIP": "🔴", "PENDING_JD": "⚪"}[rec]
            L.append(f"## {emoji} {rec}")
            L.append("")
            if rec == "PENDING_JD":
                L.append("> ⚠ 以下岗位尚未获取 JD 全文，当前评分仅基于列表页标签。")
                L.append("> 请在浏览器打开链接清单，补充 JD 后运行 `python -m job_copilot enrich --import-jd <文件>` 重新精排。")
                L.append("")

        r = s.get("_job", {})
        L.append(f"### {i+1}. {r.get('title','')} — {r.get('company','')}")
        L.append("")
        L.append(f"- **推荐**：{rec} ｜ 匹配度 {s.get('fit_score','?')}/5 ｜ "
                 f"置信度 {s.get('confidence','?')} ｜ {r.get('salary_text','')} ｜ "
                 f"{r.get('experience','')} ｜ {r.get('city','')}")
        if s.get("_jd_length"):
            L.append(f"- **JD**：已获取（{s['_jd_length']}字）")
        else:
            L.append(f"- **JD**：⚠ 未获取，等待手动补采")
        if s.get("score_change_note"):
            L.append(f"- **评分变化**：{s['score_change_note']}")
        L.append(f"- **理由**：{'；'.join(s.get('reasons',[]))}")
        if s.get("matched_evidence"):
            L.append(f"- **匹配证据**：")
            for ev in s["matched_evidence"]:
                L.append(f"  - {ev}")
        if s.get("hard_blockers"):
            L.append(f"- **硬性阻塞**：{'；'.join(s['hard_blockers'])}")
        if s.get("acceptable_gaps"):
            L.append(f"- **可接受缺口**：{'；'.join(s['acceptable_gaps'])}")
        L.append(f"- **简历突出**：{'；'.join(s.get('highlights',[]))}")
        L.append(f"- [职位链接]({r.get('url','')})")
        L.append("")

    L.append("---")
    L.append(f"*数据限制：{fr_success}/{len(candidates)} 条 JD 成功补采。"
             f"Boss直聘列表页标签 vs 详情页全文可能存在信息差。*")

    md_path = REPORTS_DIR / f"投递候选清单-{today}.md"
    md_path.write_text("\n".join(L), encoding="utf-8")

    # --- CSV ---
    csv_path = ENRICH_DIR / "top30_jobs.csv"
    csv_fields = ["recommendation", "fit_score", "verdict", "archetype", "confidence",
                  "title", "company", "salary_text", "experience", "degree", "city",
                  "industry", "company_size", "matched_evidence", "hard_blockers",
                  "acceptable_gaps", "reasons", "highlights", "score_change_note",
                  "jd_length", "jd_fetched", "url"]
    ENRICH_DIR.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=csv_fields, extrasaction="ignore")
        w.writeheader()
        for s in enriched:
            r = s.get("_job", {})
            w.writerow({
                "recommendation": s.get("recommendation"),
                "fit_score": s.get("fit_score"),
                "verdict": s.get("verdict"),
                "archetype": s.get("archetype"),
                "confidence": s.get("confidence"),
                "title": r.get("title"),
                "company": r.get("company"),
                "salary_text": r.get("salary_text"),
                "experience": r.get("experience"),
                "degree": r.get("degree"),
                "city": r.get("city"),
                "industry": r.get("industry"),
                "company_size": r.get("company_size"),
                "matched_evidence": " | ".join(s.get("matched_evidence", [])),
                "hard_blockers": " | ".join(s.get("hard_blockers", [])),
                "acceptable_gaps": " | ".join(s.get("acceptable_gaps", [])),
                "reasons": "; ".join(s.get("reasons", [])),
                "highlights": "; ".join(s.get("highlights", [])),
                "score_change_note": s.get("score_change_note"),
                "jd_length": s.get("_jd_length", 0),
                "jd_fetched": s.get("_fetch_success", False),
                "url": r.get("url"),
            })

    # --- Excel ---
    xlsx_path = ENRICH_DIR / "top30_jobs.xlsx"
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "投递候选"

    xl_headers = ["推荐", "匹配分", "画像", "置信度", "职位", "公司", "薪资",
                  "经验", "学历", "城市", "行业", "规模",
                  "匹配证据", "硬性阻塞", "可接受缺口", "评分变化", "JD字数", "链接"]
    header_font = Font(bold=True, size=11)
    hdr_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    grey_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    rec_fills = {"APPLY_NOW": green_fill, "REVIEW": yellow_fill, "SKIP": red_fill, "PENDING_JD": grey_fill}

    for ci, h in enumerate(xl_headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = hdr_fill

    for ri, s in enumerate(enriched, 2):
        r = s.get("_job", {})
        row_data = [
            s.get("recommendation"), s.get("fit_score"), s.get("archetype"),
            s.get("confidence"), r.get("title"), r.get("company"),
            r.get("salary_text"), r.get("experience"), r.get("degree"),
            r.get("city"), r.get("industry"), r.get("company_size"),
            " | ".join(s.get("matched_evidence", [])),
            " | ".join(s.get("hard_blockers", [])),
            " | ".join(s.get("acceptable_gaps", [])),
            s.get("score_change_note"),
            s.get("_jd_length", 0),
            r.get("url"),
        ]
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
        # 推荐列着色
        rec = s.get("recommendation", "")
        if rec in rec_fills:
            ws.cell(row=ri, column=1).fill = rec_fills[rec]

    widths = [12, 8, 14, 8, 28, 20, 12, 8, 6, 6, 14, 12, 40, 30, 30, 35, 7, 36]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(xlsx_path)

    return md_path, csv_path, xlsx_path


def import_jd(path: str) -> int:
    """从 JSON 文件导入 JD 文本：{"job_id": "JD全文", ...}。返回导入数。"""
    p = Path(path).expanduser()
    if not p.exists():
        print(f"✗ 文件不存在: {p}")
        return 0
    data = json.loads(p.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        print("✗ JSON 格式错误：需要 {\"job_id\": \"JD全文\", ...}")
        return 0

    conn = db.connect()
    count = 0
    for job_id, jd_text in data.items():
        if jd_text and len(str(jd_text).strip()) > 20:
            conn.execute(
                "UPDATE jobs SET jd_text = ? WHERE job_id = ? AND (jd_text IS NULL OR jd_text = '')",
                (str(jd_text).strip(), job_id),
            )
            if conn.total_changes > 0:
                count += 1
    conn.commit()
    print(f"✓ 已导入 {count} 条 JD 全文（共 {len(data)} 个 job_id）")
    return count


def _write_url_list(candidates: list[dict]) -> Path:
    """生成 Top N 岗位链接清单 HTML，方便用户在浏览器中逐条打开。"""
    today = date.today().isoformat()
    rows_html = ""
    for i, r in enumerate(candidates, 1):
        url = r.get("url", "")
        rows_html += (
            f'<tr>'
            f'<td>{i}</td>'
            f'<td><a href="{url}" target="_blank">{r.get("title","")}</a></td>'
            f'<td>{r.get("company","")}</td>'
            f'<td>{r.get("salary_text","")}</td>'
            f'<td>{r.get("experience","")}</td>'
            f'<td>{r.get("city","")}</td>'
            f'<td class="jd" data-id="{r.get("job_id","")}" contenteditable="true"></td>'
            f'</tr>'
        )

    html = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Top {len(candidates)} 岗位 JD 补采清单</title>
<style>
body {{ font-family: system-ui, sans-serif; max-width: 1200px; margin: 0 auto; padding: 20px; background: #f5f5f5; }}
h1 {{ font-size: 20px; }}
table {{ width: 100%; border-collapse: collapse; background: white; border-radius: 8px; overflow: hidden; }}
th, td {{ padding: 8px 10px; border-bottom: 1px solid #e0e0e0; font-size: 13px; text-align: left; }}
th {{ background: #2a78d6; color: white; font-weight: 600; }}
td.jd {{ min-width: 300px; max-width: 500px; font-size: 12px; color: #666; white-space: pre-wrap; }}
a {{ color: #2a78d6; text-decoration: none; }}
a:hover {{ text-decoration: underline; }}
.help {{ background: #fff3cd; border: 1px solid #ffc107; border-radius: 8px; padding: 12px 16px; margin: 16px 0; font-size: 13px; }}
button {{ padding: 8px 16px; background: #2a78d6; color: white; border: none; border-radius: 6px; cursor: pointer; font-size: 13px; }}
button:hover {{ background: #1a5db8; }}
#status {{ font-size: 12px; color: #666; margin-left: 8px; }}
</style>
</head>
<body>
<h1>Top {len(candidates)} 岗位 · JD 补采清单</h1>
<div class="help">
  <strong>使用方法：</strong>
  <br>1. 点击岗位链接，在新标签页打开 Boss直聘 详情页
  <br>2. 复制岗位描述文本（岗位职责 + 任职要求）
  <br>3. 粘贴到对应行的「JD 文本」列
  <br>4. 全部填完后点击下方「导出 JSON」按钮
  <br>5. 在终端运行：<code>python -m job_copilot enrich --import-jd ~/Downloads/jd_export.json</code>
</div>
<table>
<thead><tr>
  <th>#</th><th>职位</th><th>公司</th><th>薪资</th><th>经验</th><th>城市</th><th>JD 文本（粘贴到这里）</th>
</tr></thead>
<tbody>{rows_html}</tbody>
</table>
<div style="margin-top: 16px;">
  <button onclick="exportJSON()">📋 导出 JSON</button>
  <span id="status"></span>
</div>
<script>
function exportJSON() {{
  const rows = document.querySelectorAll('td.jd');
  const data = {{}};
  rows.forEach(td => {{
    const jd = td.innerText.trim();
    if (jd) data[td.dataset.id] = jd;
  }});
  if (Object.keys(data).length === 0) {{
    document.getElementById('status').textContent = '请先在 JD 列粘贴内容';
    return;
  }}
  const blob = new Blob([JSON.stringify(data, null, 2)], {{type: 'application/json'}});
  const a = document.createElement('a');
  a.href = URL.createObjectURL(blob);
  a.download = 'jd_export.json';
  a.click();
  document.getElementById('status').textContent = `已导出 ${{Object.keys(data).length}} 条 JD`;
}}
</script>
</body>
</html>"""

    path = ENRICH_DIR / f"jd_fetch_list-{today}.html"
    ENRICH_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path


# ============================================================
# CLI 入口
# ============================================================

def run(top: int = 30, fetch: bool = True, rescore: bool = True,
        import_jd_path: str = "") -> str:
    # 0. JD 导入模式
    if import_jd_path:
        n = import_jd(import_jd_path)
        if n == 0:
            return "✗ 未能导入任何 JD。"

    # 1. 选择 Top N
    candidates = select_top(n=top)
    if not candidates:
        return "✗ 无已精排岗位可选。先运行 score。"

    print(f"→ 从 {_count_scored()} 个已精排岗位中选出 Top {len(candidates)}")

    # 生成链接清单 HTML（无论是否 fetch，都生成方便手动补采）
    url_html = _write_url_list(candidates)

    # 2. JD 补采
    if fetch:
        fetch_results = fetch_details(candidates)
    else:
        fetch_results = [FetchResult(
            job_pk=r["id"], job_id=r.get("job_id", ""),
            url=r.get("url", ""), success=bool(r.get("jd_text")),
            jd_text=r.get("jd_text", ""),
            failure_reason="" if r.get("jd_text") else "跳过补采",
        ) for r in candidates]

    # 3. 二次精排
    if rescore:
        enriched = rescore_all(candidates, fetch_results)
    else:
        enriched = []
        conn = db.connect()
        resume = score_mod._resume_text()
        for r in candidates:
            s = rescore_one(conn, r, resume, "")
            s["_job"] = r
            s["_fetch_success"] = bool(r.get("jd_text"))
            s["_jd_length"] = len(r.get("jd_text") or "")
            enriched.append(s)

    # 4. 报告
    md_path, csv_path, xlsx_path = _write_reports(enriched, candidates, fetch_results)

    # 统计
    fr_ok = sum(1 for fr in fetch_results if fr.success)
    fr_fail = len(fetch_results) - fr_ok
    rec = Counter(s.get("recommendation") for s in enriched)

    failure_reasons = Counter(fr.failure_reason for fr in fetch_results if not fr.success)
    fail_summary = "、".join(f"{k}({v})" for k, v in failure_reasons.most_common(3)) if failure_reasons else "无"

    jd_count = sum(1 for s in enriched if s.get("_jd_length", 0) > 0)
    pending_count = rec.get("PENDING_JD", 0)
    title = "JD 补采 + 二次精排完成" if jd_count else "候选清单已生成（待补充 JD）"

    lines = [
        "=" * 55,
        f"  {title}",
        "=" * 55,
        f"  候选岗位数       : {len(candidates)}",
        f"  JD 补采成功      : {fr_ok}",
        f"  JD 补采失败      : {fr_fail}（{fail_summary}）",
        f"  二次精排数       : {len(enriched)}",
        f"  APPLY_NOW       : {rec.get('APPLY_NOW', 0)}",
        f"  REVIEW          : {rec.get('REVIEW', 0)}",
        f"  SKIP            : {rec.get('SKIP', 0)}",
        f"  PENDING_JD      : {pending_count}",
        "",
        f"  链接清单(手动补采) : {url_html}",
        f"  候选清单 Markdown : {md_path}",
        f"  候选清单 CSV      : {csv_path}",
        f"  候选清单 Excel    : {xlsx_path}",
    ]

    if pending_count > 0:
        lines += [
            "",
            f"  💡 {pending_count} 个岗位待补充 JD：用浏览器打开链接清单，"
            f"粘贴 JD → 导出JSON → ",
            f"     python -m job_copilot enrich --import-jd <文件>",
        ]
    lines.append("=" * 55)
    return "\n".join(lines)


def _count_scored() -> int:
    conn = db.connect()
    return conn.execute(
        "SELECT COUNT(*) FROM job_scores s JOIN jobs j ON j.id=s.job_pk "
        "WHERE s.seniority_ok=1 AND s.authenticity!='低' AND j.url IS NOT NULL AND j.url!=''"
    ).fetchone()[0]
